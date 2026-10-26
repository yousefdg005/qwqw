import base64
import hashlib
import io
import json
import os
import re
import time
import unicodedata
import xml.etree.ElementTree as ET
from html import escape as html_escape
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

import arabic_reshaper
import pandas as pd
from bidi.algorithm import get_display
from docx import Document
from openai import OpenAI
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
import streamlit as st
import streamlit.components.v1 as components
import requests

def send_to_dictionary(extracted_terms):
    url = "https://dictionaryly.onrender.com/api/receive-terms"
    headers = {
        "Authorization": "Bearer my_secure_platform_ingest_token_2026",
        "Content-Type": "application/json"
    }
    
    # تقسيم المصطلحات إلى دفعات (كل دفعة 10 مصطلحات) لتجنب الضغط وخطأ 500
    batch_size = 10
    success_all = True
    
    for i in range(0, len(extracted_terms), batch_size):
        batch = extracted_terms[i:i + batch_size]
        payload = {"terms": batch}
        try:
            response = requests.post(url, json=payload, headers=headers, timeout=60)
            if response.status_code != 201:
                success_all = False
                st.error(f"❌ تم الرفض في الدفعة {i//batch_size + 1}: كود الخطأ {response.status_code} - التفاصيل: {response.text}")
        except Exception as e:
            success_all = False
            st.error(f"❌ فشل الاتصال في الدفعة {i//batch_size + 1}: {e}")
            
    return success_all

# =========================================================
# 1. إعدادات صفحة Streamlit
# =========================================================
st.set_page_config(
    page_title="AI Legal Workspace",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# =========================================================
# 2. تحميل ملف الـ CSS
# =========================================================
def load_css(file_name="style.css"):
    try:
        css_path = Path(__file__).parent / file_name
        with open(css_path, "r", encoding="utf-8") as f:
            st.markdown(f"<style>{f.read()}</style>", unsafe_allow_html=True)
    except FileNotFoundError:
        pass

load_css("style.css")

# =========================================================
# 3. دالة معالجة النص العربي للـ PDF (ReportLab)
# =========================================================
def prepare_arabic_text(text):
    if not text or not isinstance(text, str):
        return ""
    reshaped_text = arabic_reshaper.reshape(text)
    bidi_text = get_display(reshaped_text)
    return bidi_text

# =========================================================
# 4. الهيدر العام للمنصة (Neutral Branding)
# =========================================================
st.markdown(
    """
<div class="header-container" style="direction: ltr; text-align: left;">
    <div class="header-text">
        <h1 style="color: #EAB308; margin: 0; font-size: 28px; font-weight: 900; letter-spacing: -0.5px;">
            ⚖️ AI Legal Workspace
        </h1>
        <p style="color: #8B9DC3; font-size: 14px; margin: 10px 0 14px; line-height: 1.6;">
            Your all-in-one AI toolkit for legal text alignment &amp; term extraction
        </p>
        <div style="display: flex; gap: 10px; justify-content: center; flex-wrap: wrap;">
            <span style="color: #EEF2F8; font-size: 12px; background: rgba(234,179,8,0.12); padding: 5px 14px; border-radius: 20px; border: 1px solid rgba(234,179,8,0.3); font-weight: 600;">
                🤖 Smart AI Engine
            </span>
            <span style="color: #EEF2F8; font-size: 12px; background: rgba(59,130,246,0.1); padding: 5px 14px; border-radius: 20px; border: 1px solid rgba(59,130,246,0.25); font-weight: 600;">
                🌐 AR ↔ EN
            </span>
            <span style="color: #8B9DC3; font-size: 12px; background: rgba(139,157,195,0.08); padding: 5px 14px; border-radius: 20px; border: 1px solid rgba(139,157,195,0.18);">
                💻 Developed by Youssef Nour Eldin
            </span>
        </div>
    </div>
</div>
""",
    unsafe_allow_html=True,
)

# =========================================================
# 5. الشريط الجانبي للإعدادات
# =========================================================
st.sidebar.header("⚙️ إعدادات النظام")

st.sidebar.markdown(
    """
<div style="background-color: #1E293B; padding: 12px; border-radius: 8px; border: 1px solid #374151; margin-bottom: 15px; text-align: center;">
    <span style="color: #E5E7EB; font-size: 13px;">ليس لديك مفتاح API؟</span><br>
    <a href="https://openrouter.ai/keys" target="_blank" style="color: #EAB308; font-weight: bold; text-decoration: none; font-size: 14px; display: inline-block; margin-top: 5px;">
        🔗 الحصول على المفتاح من OpenRouter
    </a>
</div>
""",
    unsafe_allow_html=True,
)

api_key = st.sidebar.text_input(
    "🔑 مفتاح OpenRouter API Key الخاص بك:",
    value="",
    type="password",
    placeholder="sk-or-v1-...",
)

PAID_MODELS_MAPPING = {
    "⚡ Gemini 3.8 Flash": "google/gemini-3.8-flash",
    "⚖️ GPT-4o Mini": "openai/gpt-4o-mini",
    "🦙 Llama 4 Maverick": "meta-llama/llama-4-maverick",
}

st.sidebar.subheader("🤖 اختيار النموذج المراد استعماله")
selected_model_label = st.sidebar.selectbox(
    "اختر النموذج لتنفيذ عملية التصفية والاستخراج:",
    options=list(PAID_MODELS_MAPPING.keys()),
    index=0,
)
selected_model_id = PAID_MODELS_MAPPING[selected_model_label]

with st.sidebar:
    st.markdown("---")
    st.subheader("📚 القاموس القانوني")
    st.link_button("🌐 فتح القاموس", "https://dictionaryly.onrender.com", use_container_width=True)

# =========================================================
# 6. إدارة حالة الجلسة (Session State)
# =========================================================
if "aligned_pairs" not in st.session_state:
    st.session_state["aligned_pairs"] = []
if "results_df" not in st.session_state:
    st.session_state["results_df"] = None
if "tmx_bytes" not in st.session_state:
    st.session_state["tmx_bytes"] = None
if "alignment_draft" not in st.session_state:
    st.session_state["alignment_draft"] = None


# =========================================================
# 7. دوال استخراج وقراءة ملفات Word و TMX
# =========================================================
_INVISIBLE_RE = re.compile(r"[\u200b-\u200f\u202a-\u202e\u2066-\u2069\ufeff]")
_XML_BAD_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

def normalize_text(text):
    text = unicodedata.normalize("NFKC", text)
    text = _INVISIBLE_RE.sub("", text)
    return text.replace("\u00a0", " ")

_EN_STOP = set("a an the and or but if then of to in on at by for from with without under over into upon is are was were be been being it this that these those as not no any all each such may shall will must can could should has have had do does did".split())
_AR_STOP = set("في من على إلى عن أن إن أو و ثم كما حيث هذا هذه ذلك التي الذي الذين اللذين كان تكون يكون يتم قد ما لا لم لن مع بين عند لدى كل أي بعض وفق بموجب بناء بعد قبل حيث إذا".split())

def _term_key(value):
    value = normalize_text(str(value or "")).casefold()
    value = re.sub(r"[\W_]+", " ", value, flags=re.UNICODE)
    return re.sub(r"\s+", " ", value).strip()

def align_segments(en_segments, ar_segments):
    n, m = len(en_segments), len(ar_segments)
    if not n or not m:
        count = max(n, m)
        return [(en_segments[i] if i < n else "", ar_segments[i] if i < m else "")
                for i in range(count)]
    inf = float("inf")
    dp = [[inf] * (m + 1) for _ in range(n + 1)]
    prev = [[None] * (m + 1) for _ in range(n + 1)]
    dp[0][0] = 0.0
    for i in range(n + 1):
        for j in range(m + 1):
            if dp[i][j] == inf: continue
            for a, b in ((1,1),(1,2),(2,1)):
                if i+a > n or j+b > m: continue
                le = sum(len(x) for x in en_segments[i:i+a]); la = sum(len(x) for x in ar_segments[j:j+b])
                cost = abs(__import__('math').log((le+1)/(la+1))) * 0.10 + (a+b-2)*0.18
                if dp[i][j] + cost < dp[i+a][j+b]:
                    dp[i+a][j+b] = dp[i][j] + cost; prev[i+a][j+b] = (i,j,a,b)
    if prev[n][m] is None:
        return [(en_segments[i] if i<n else "", ar_segments[i] if i<m else "") for i in range(max(n,m))]
    groups=[]; i,j=n,m
    while i or j:
        step=prev[i][j]
        if step is None: break
        pi,pj,a,b=step
        groups.append((" ".join(en_segments[pi:i]), " ".join(ar_segments[pj:j])))
        i,j=pi,pj
    return list(reversed(groups))

def _is_low_value_term(term):
    key = _term_key(term)
    if not key: return True
    words = key.split()
    if len(words) == 1 and (key in _EN_STOP or key in _AR_STOP): return True
    if len(words) == 1 and len(key) < 4 and not re.search(r"[\u0600-\u06ff]", key): return True
    return False

def extract_text_from_docx_bytes(data):
    notes = []
    try:
        doc = Document(io.BytesIO(data))
        paragraphs = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                paragraphs.append(normalize_text(txt))
        if doc.tables:
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        paragraphs.append(normalize_text(" | ".join(row_cells)))
            notes.append(f"تم استخراج المحتوى من {len(doc.tables)} جدول/جداول داخل المستند")
        return "\n\n".join(paragraphs), notes
    except Exception as e:
        return "", [f"خطأ أثناء قراءة ملف Word: {e}"]

def process_docx_text_to_segments(text):
    if not text: return []
    text = normalize_text(text)
    segments = []
    abbreviations = ("Mr.", "Mrs.", "Ms.", "Dr.", "Jr.", "Sr.", "Inc.", "Ltd.", "No.", "e.g.", "i.e.", "Art.", "Sec.")
    for paragraph in re.split(r"\n+", text):
        paragraph = re.sub(r"\s+", " ", paragraph).strip()
        if not paragraph: continue
        protected = paragraph
        for idx, abbr in enumerate(abbreviations):
            protected = re.sub(re.escape(abbr), f"{abbr[:-1]}<ABBR{idx}>", protected, flags=re.I)
        parts = re.split(r"(?<=[.!?؟؛])\s+", protected)
        for part in parts:
            part = re.sub(r"<ABBR(\d+)>", lambda m: abbreviations[int(m.group(1))][-1], part)
            part = re.sub(r"\s+", " ", part).strip()
            if part: segments.append(part)
    return segments

def parse_tmx_bytes(tmx_content):
    try:
        root = ET.fromstring(tmx_content)
    except (ET.ParseError, TypeError, ValueError) as exc:
        raise ValueError(f"ملف TMX غير صالح: {exc}") from exc
    pairs=[]
    for tu in root.iter():
        if tu.tag.split("}")[-1] != "tu": continue
        langs={}
        for tuv in list(tu):
            if tuv.tag.split("}")[-1] != "tuv": continue
            lang=next((v.lower() for k,v in tuv.attrib.items() if k.lower().endswith("lang")), "")
            seg=next((el for el in tuv.iter() if el.tag.split("}")[-1]=="seg"), None)
            value="".join(seg.itertext()).strip() if seg is not None else ""
            if lang.startswith("en"): langs["en"]=value
            elif lang.startswith("ar"): langs["ar"]=value
        if langs.get("en") or langs.get("ar"): pairs.append((langs.get("en",""),langs.get("ar","")))
    return pairs

def generate_tmx_file(pairs):
    root=ET.Element("tmx", {"version":"1.4"})
    ET.SubElement(root,"header", {"creationtool":"AI Legal Workspace","creationtoolversion":"3.0","segtype":"sentence","adminlang":"en-us","srclang":"en","datatype":"plaintext"})
    body=ET.SubElement(root,"body")
    for en, ar in pairs:
        tu=ET.SubElement(body,"tu")
        for lang,text in (("en",en),("ar",ar)):
            tuv=ET.SubElement(tu,"tuv", {"{http://www.w3.org/XML/1998/namespace}lang":lang})
            ET.SubElement(tuv,"seg").text=_XML_BAD_RE.sub("",str(text or ""))
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


# =========================================================
# 8. دالة توليد بطاقات الـ PDF
# =========================================================
_FONT_CANDIDATES = [
    ("fonts/Amiri-Regular.ttf", "fonts/Amiri-Bold.ttf"),
    ("fonts/NotoNaskhArabic-Regular.ttf", "fonts/NotoNaskhArabic-Bold.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf", "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Bold.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("C:/Windows/Fonts/tahoma.ttf", "C:/Windows/Fonts/tahomabd.ttf"),
    ("/Library/Fonts/Arial Unicode.ttf", "/Library/Fonts/Arial Unicode.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
]
_PDF_FONTS = None

def get_pdf_fonts():
    global _PDF_FONTS
    if _PDF_FONTS:
        return _PDF_FONTS
    for reg, bold in _FONT_CANDIDATES:
        if os.path.exists(reg):
            try:
                pdfmetrics.registerFont(TTFont("AppFont", reg))
                pdfmetrics.registerFont(TTFont("AppFont-Bold", bold if os.path.exists(bold) else reg))
                _PDF_FONTS = ("AppFont", "AppFont-Bold")
                return _PDF_FONTS
            except Exception:
                continue
    _PDF_FONTS = ("Helvetica", "Helvetica-Bold")
    return _PDF_FONTS

def _wrap_logical(text, font, size, max_width):
    lines, cur = [], ""
    for word in str(text).split():
        cand = f"{cur} {word}".strip()
        if cur and pdfmetrics.stringWidth(prepare_arabic_text(cand), font, size) > max_width:
            lines.append(cur)
            cur = word
        else:
            cur = cand
    if cur:
        lines.append(cur)
    return lines or [""]

def ar_par(text, style, max_width):
    lines = _wrap_logical(text, style.fontName, style.fontSize, max_width)
    html = "<br/>".join(html_escape(prepare_arabic_text(l)) for l in lines)
    return Paragraph(html, style)

def en_par(text, style):
    return Paragraph(html_escape(str(text)), style)

def generate_pdf_cards_bytes(df):
    font, font_bold = get_pdf_fonts()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    story = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle("TitleStyle", parent=styles["Heading1"], fontName=font_bold, fontSize=18, alignment=1, textColor=colors.HexColor("#0F172A"), leading=26)
    sub_style = ParagraphStyle("SubStyle", parent=styles["Normal"], fontName=font, fontSize=10, alignment=1, textColor=colors.HexColor("#64748B"), leading=14)
    cell_title_style = ParagraphStyle("CellTitle", fontName=font_bold, fontSize=10, alignment=2, leading=15, textColor=colors.HexColor("#1E293B"))
    cell_val_style = ParagraphStyle("CellVal", fontName=font, fontSize=10, alignment=2, leading=15, textColor=colors.HexColor("#0F172A"))
    cell_val_en_style = ParagraphStyle("CellValEn", parent=cell_val_style, alignment=0)

    VAL_W, LBL_W, PAD = 360, 160, 6
    val_text_w = VAL_W - 2 * PAD - 4
    lbl_text_w = LBL_W - 2 * PAD - 4

    story.append(ar_par("بطاقات المصطلحات القانونية الموثقة (Legal Terminology Cards)", title_style, 520))
    story.append(ar_par("منصة التحليل والاستخراج القانوني الذكي | إعداد: يوسف نورالدين", sub_style, 520))
    story.append(Spacer(1, 15))

    total = len(df)
    for n, (_, row) in enumerate(df.iterrows(), start=1):
        story.append(ar_par(f"بطاقة مصطلح رقم: ({n:02d} / {total:02d})", cell_title_style, 520))
        story.append(Spacer(1, 5))

        def lbl(t): return ar_par(t, cell_title_style, lbl_text_w)
        def ar_val(key, default=""): return ar_par(str(row.get(key, default) or default), cell_val_style, val_text_w)

        data = [
            [en_par(row.get("en", ""), cell_val_en_style), lbl("المصطلح بالإنجليزي (en)")],
            [ar_val("ar"), lbl("المقابل بالعربي (ar)")],
            [ar_val("subject field", "قانون عام"), lbl("المجال (subject field)")],
            [en_par(row.get("Source EN", ""), cell_val_en_style), lbl("السياق الإنجليزي (Context EN)")],
            [ar_val("Source AR"), lbl("السياق العربي (Context AR)")],
            [ar_val("definition"), lbl("الشرح والتعريف (definition)")],
        ]

        t = Table(data, colWidths=[VAL_W, LBL_W])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("PADDING", (0, 0), (-1, -1), PAD),
        ]))
        story.append(t)
        story.append(Spacer(1, 15))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#E2E8F0"), spaceAfter=15))

    doc.build(story)
    buffer.seek(0)
    return buffer

# =========================================================
# 9. علامات التبويب الرئيسية (Tabs)
# =========================================================
tab1, tab2 = st.tabs(["Alignment Studio (beta)", "AI Term Extractor"])

# =========================================================
# مكوّن استوديو المحاذاة
# =========================================================
ALIGNMENT_STUDIO_HTML = r"""<!DOCTYPE html>
<html lang="ar"><head><meta charset="utf-8">
<style>
  * { box-sizing: border-box; }
  html, body { height: 100%; overflow: hidden; }
  body { font-family: 'Cairo','Segoe UI',Tahoma,Arial,sans-serif; background:#07090F; color:#EEF2F8; margin:0; padding:10px; overflow:hidden; }
  #bar { display:flex; gap:10px; align-items:center; margin-bottom:12px; direction:rtl; flex-wrap:wrap; }
  #toast { color:#FDE047; font-size:13px; font-weight:600; }
  #alignmentScroll {
    height: 610px; max-height: calc(100vh - 140px); min-height: 240px;
    overflow-y: auto; overflow-x: hidden; padding: 10px;
    border: 1px solid #1B2840; border-radius: 14px; background: #0D1424;
    scrollbar-width: thin; scrollbar-color: #243350 #0D1424; overscroll-behavior: contain;
  }
  #alignmentScroll::-webkit-scrollbar { width: 7px; }
  #alignmentScroll::-webkit-scrollbar-track { background: #0D1424; border-radius: 4px; }
  #alignmentScroll::-webkit-scrollbar-thumb { background: #243350; border-radius: 4px; }
  #alignmentScroll::-webkit-scrollbar-thumb:hover { background: #344d6e; }
  .alignment-container { display:flex; flex-direction:column; gap:10px; direction:rtl; padding-bottom:14px; }
  .segment-row { display:flex; gap:10px; background:#0F1D35; border:1px solid #1B2840; border-radius:12px; padding:12px 14px; align-items:center; transition: border-color 0.2s, box-shadow 0.2s; position: relative; }
  .segment-row::before { content:''; position:absolute; right:0; top:8px; bottom:8px; width:3px; background:linear-gradient(180deg,#EAB308,#CA8A04); border-radius:3px; opacity:0; transition:opacity 0.2s; }
  .segment-row:hover { border-color:#243350; box-shadow:0 2px 16px rgba(0,0,0,0.35); }
  .segment-row:hover::before { opacity:1; }
  .segment-num { color:#EAB308; font-weight:800; width:42px; text-align:center; font-size:13px; flex-shrink:0; letter-spacing:0.5px; }
  .cell { flex:1; display:flex; gap:6px; align-items:stretch; min-width:0; border:2px solid transparent; border-radius:8px; transition:border-color 0.18s, background 0.18s; }
  .cell.drag-over { border:2px dashed #EAB308; background:rgba(234,179,8,0.1); }
  .handle-container { display:flex; flex-direction:column; gap:2px; flex-shrink:0; align-items:center; justify-content:center; }
  .handle { width:22px; height:22px; display:flex; align-items:center; justify-content:center; background:#111D33; border-radius:5px; color:#4D6080; cursor:grab; user-select:none; font-size:12px; border:1px solid #1B2840; transition:all 0.18s; }
  .handle:hover { color:#EAB308; border-color:rgba(234,179,8,0.3); background:#182640; }
  .move-btn { width:22px; height:18px; display:flex; align-items:center; justify-content:center; background:#111D33; border:1px solid #1B2840; border-radius:4px; color:#4D6080; cursor:pointer; font-size:10px; padding:0; transition:all 0.18s; }
  .move-btn:hover { background:#182640; color:#EAB308; border-color:rgba(234,179,8,0.3); }
  .card-box { flex:1; min-width:0; background:#111D33; border:1px solid #1B2840; border-radius:8px; padding:10px 14px; min-height:55px; font-size:14px; line-height:1.75; white-space:pre-wrap; overflow-wrap:anywhere; cursor:text; transition:border-color 0.18s, background 0.18s, box-shadow 0.18s; }
  .card-box:focus { outline:none; border-color:#EAB308; background:#0D1424; box-shadow:0 0 0 3px rgba(234,179,8,0.1); }
  .en-box { direction:ltr; text-align:left; color:#7DBAFF; }
  .ar-box { direction:rtl; text-align:right; color:#FDE047; }
  .action-btn-group { display:flex; flex-direction:column; gap:5px; flex-shrink:0; }
  .btn-ctrl { background:#111D33; color:#8B9DC3; border:1px solid #1B2840; border-radius:6px; padding:5px 9px; font-size:11px; font-weight:700; cursor:pointer; font-family:inherit; transition:all 0.18s; }
  .btn-ctrl:hover { background:#182640; color:#EAB308; border-color:rgba(234,179,8,0.3); }
  .btn-del:hover { background:#3B0A0A; color:#FCA5A5; border-color:rgba(239,68,68,0.3); }
  .add-row-btn { background:linear-gradient(135deg,#EAB308 0%,#CA8A04 100%); color:#07090F; font-weight:800; border:none; border-radius:8px; padding:10px 20px; cursor:pointer; font-family:inherit; font-size:13px; box-shadow:0 4px 14px rgba(234,179,8,0.28); transition:all 0.2s; }
  .add-row-btn:hover { transform:translateY(-2px); box-shadow:0 6px 22px rgba(234,179,8,0.4); filter:brightness(1.06); }
  .undo-btn { background:#111D33; color:#8B9DC3; border:1px solid #1B2840; border-radius:8px; padding:10px 16px; cursor:pointer; font-weight:700; font-family:inherit; font-size:13px; transition:all 0.18s; }
  .undo-btn:hover { background:#182640; color:#EEF2F8; border-color:#243350; }
</style></head>
<body>
<div id="bar">
  <button class="add-row-btn" id="addBtn">➕ إضافة وحدة محاذاة جديدة</button>
  <button class="undo-btn" id="undoBtn">↩️ تراجع</button>
  <span id="toast"></span>
</div>
<div id="alignmentScroll" aria-label="منطقة تمرير وحدات المحاذاة"><div class="alignment-container" id="alignmentContainer"></div></div>
<script>
  const container = document.getElementById('alignmentContainer');
  const toastEl = document.getElementById('toast');
  let data = [], historyStack = [], token = null, drag = null, sendTimer = null;

  function post(type, extra) { window.parent.postMessage(Object.assign({isStreamlitMessage:true, type}, extra||{}), '*'); }
  function pushValue() { post('streamlit:setComponentValue', {value: data.map(d => ({en:d.en, ar:d.ar})), dataType:'json'}); }
  function pushSoon() { clearTimeout(sendTimer); sendTimer = setTimeout(pushValue, 500); }
  
  function saveState() {
    historyStack.push(JSON.stringify(data));
    if (historyStack.length > 50) historyStack.shift();
  }

  window.addEventListener('message', (ev) => {
    if (!ev.data || ev.data.type !== 'streamlit:render') return;
    const args = ev.data.args || {};
    if (args.token !== token) {
      token = args.token;
      data = (args.pairs || []).map(p => ({en: p.en || '', ar: p.ar || ''}));
      historyStack = [];
      renderRows();
    }
  });
  post('streamlit:componentReady', {apiVersion: 1});
  post('streamlit:setFrameHeight', {height: 800});

  function toast(msg) { toastEl.textContent = msg; setTimeout(() => { if (toastEl.textContent === msg) toastEl.textContent = ''; }, 2500); }

  function moveCell(index, type, direction) {
    const targetIndex = index + direction;
    if (targetIndex < 0 || targetIndex >= data.length) return;
    saveState();
    const temp = data[index][type];
    data[index][type] = data[targetIndex][type];
    data[targetIndex][type] = temp;
    renderRows(); pushValue();
  }

  function makeCell(type, index) {
    const cell = document.createElement('div');
    cell.className = 'cell'; cell.dataset.index = index; cell.dataset.type = type;
    
    const handleContainer = document.createElement('div');
    handleContainer.className = 'handle-container';
    
    const upBtn = document.createElement('button');
    upBtn.className = 'move-btn'; upBtn.textContent = '▲'; upBtn.title = 'نقل للأعلى';
    upBtn.onclick = () => moveCell(index, type, -1);
    
    const handle = document.createElement('div');
    handle.className = 'handle'; handle.textContent = '⠿'; handle.draggable = true; handle.title = 'اسحب للتبديل';
    
    const downBtn = document.createElement('button');
    downBtn.className = 'move-btn'; downBtn.textContent = '▼'; downBtn.title = 'نقل للأسفل';
    downBtn.onclick = () => moveCell(index, type, 1);
    
    handleContainer.append(upBtn, handle, downBtn);

    const box = document.createElement('div');
    box.className = 'card-box ' + (type === 'en' ? 'en-box' : 'ar-box');
    box.dataset.index = index; box.dataset.type = type;
    box.contentEditable = 'plaintext-only';
    if (box.contentEditable !== 'plaintext-only') box.contentEditable = 'true';
    box.textContent = data[index][type];
    
    let prevText = box.textContent;
    box.addEventListener('focus', () => { prevText = box.innerText; });
    box.addEventListener('blur', () => {
      const currentText = box.innerText.replace(/\n$/, '');
      if (prevText !== currentText) { saveState(); data[index][type] = currentText; pushSoon(); }
    });
    box.addEventListener('input', () => { data[index][type] = box.innerText.replace(/\n$/, ''); pushSoon(); });
    box.addEventListener('paste', (e) => {
      if (box.contentEditable === 'plaintext-only') return;
      e.preventDefault();
      const text = (e.clipboardData || window.clipboardData).getData('text/plain');
      const sel = window.getSelection();
      if (sel && sel.rangeCount) {
        const range = sel.getRangeAt(0);
        range.deleteContents();
        range.insertNode(document.createTextNode(text));
        range.collapse(false);
        sel.removeAllRanges();
        sel.addRange(range);
      }
      data[Number(box.dataset.index)][box.dataset.type] = box.innerText.replace(/\n$/, '');
      pushSoon();
    });

    handle.addEventListener('dragstart', (e) => { drag = {index, type}; e.dataTransfer.effectAllowed = 'move'; e.dataTransfer.setData('text/plain', String(index)); });
    cell.addEventListener('dragover', (e) => {
      if (!drag || drag.type !== type) return;
      e.preventDefault(); cell.classList.add('drag-over');
      const scroller = document.getElementById('alignmentScroll');
      const rect = scroller.getBoundingClientRect();
      const edge = 70;
      if (e.clientY < rect.top + edge) scroller.scrollTop -= 20;
      else if (rect.bottom - e.clientY < edge) scroller.scrollTop += 20;
    });
    cell.addEventListener('dragleave', () => cell.classList.remove('drag-over'));
    cell.addEventListener('drop', (e) => {
      e.preventDefault(); cell.classList.remove('drag-over');
      if (!drag || drag.type !== type || drag.index === index) { drag = null; return; }
      saveState();
      const t = data[drag.index][type]; data[drag.index][type] = data[index][type]; data[index][type] = t;
      drag = null; renderRows(); pushValue();
    });
    
    cell.append(handleContainer, box);
    return cell;
  }

  function renderRows() {
    container.textContent = '';
    data.forEach((_, index) => {
      const row = document.createElement('div'); row.className = 'segment-row';
      const num = document.createElement('div'); num.className = 'segment-num'; num.textContent = '[' + (index + 1) + ']';
      const grp = document.createElement('div'); grp.className = 'action-btn-group';
      const sp = document.createElement('button'); sp.className = 'btn-ctrl'; sp.textContent = '✂️ تقسيم'; sp.title = 'تقسيم عند مكان المؤشر';
      const del = document.createElement('button'); del.className = 'btn-ctrl btn-del'; del.textContent = '🗑️ حذف'; del.title = 'حذف هذه الوحدة';
      [sp, del].forEach(b => b.addEventListener('mousedown', e => e.preventDefault()));
      sp.addEventListener('click', () => splitUnit(index));
      del.addEventListener('click', () => deleteRow(index));
      grp.append(sp, del);
      row.append(num, makeCell('en', index), makeCell('ar', index), grp);
      container.appendChild(row);
    });
  }

  function splitUnit(index) {
    const sel = window.getSelection();
    if (!sel.rangeCount || !sel.anchorNode) { toast('ضع المؤشر داخل النص أولاً'); return; }
    const el = sel.anchorNode.nodeType === 1 ? sel.anchorNode : sel.anchorNode.parentElement;
    const box = el && el.closest('.card-box');
    if (!box || Number(box.dataset.index) !== index) { toast('ضع المؤشر داخل نص هذه الوحدة ثم اضغط تقسيم'); return; }
    const type = box.dataset.type;
    const pre = document.createRange();
    pre.selectNodeContents(box); pre.setEnd(sel.anchorNode, sel.anchorOffset);
    const full = data[index][type];
    const offset = Math.min(pre.toString().length, full.length);
    const part1 = full.slice(0, offset).trim(), part2 = full.slice(offset).trim();
    if (!part2) { toast('المؤشر في نهاية النص، لا يوجد ما يُقسَّم'); return; }
    
    saveState();
    data[index][type] = part1;
    const nr = {en: '', ar: ''}; nr[type] = part2;
    data.splice(index + 1, 0, nr);
    renderRows(); pushValue();
  }

  function deleteRow(index) { saveState(); data.splice(index, 1); renderRows(); pushValue(); }
  
  document.getElementById('addBtn').addEventListener('click', () => {
    saveState();
    data.push({en: '', ar: ''}); renderRows(); pushValue();
    const scroller = document.getElementById('alignmentScroll');
    scroller.scrollTop = scroller.scrollHeight;
  });

  document.getElementById('undoBtn').addEventListener('click', () => {
    if (historyStack.length === 0) {
      toast('لا يوجد إجراء للتراجع عنه');
      return;
    }
    data = JSON.parse(historyStack.pop());
    renderRows(); pushValue();
  });
</script></body></html>"""

_COMPONENT_DIR = Path(__file__).parent / "alignment_studio_component"

def _declare_alignment_studio():
    _COMPONENT_DIR.mkdir(exist_ok=True)
    index = _COMPONENT_DIR / "index.html"
    if not index.exists() or index.read_text(encoding="utf-8") != ALIGNMENT_STUDIO_HTML:
        index.write_text(ALIGNMENT_STUDIO_HTML, encoding="utf-8")
    return components.declare_component("alignment_studio", path=str(_COMPONENT_DIR))

alignment_studio = _declare_alignment_studio()

# ---------------------------------------------------------
# Tab 1: استوديو المحاذاة الشامل الاحترافي
# ---------------------------------------------------------
with tab1:
    st.markdown("<style>.en-box{direction:ltr;text-align:left}.ar-box{direction:rtl;text-align:right} [data-testid='stDataEditor']{direction:rtl}</style>", unsafe_allow_html=True)
    st.subheader("📄 الخطوة 1: رفع ملفات Word (.docx) واستخراج النصوص")

    col_doc1, col_doc2 = st.columns(2)
    with col_doc1:
        doc_en = st.file_uploader("ارفع ملف Word المصدر (إنجليزي):", type=["docx"], key="doc_en")
    with col_doc2:
        doc_ar = st.file_uploader("ارفع ملف Word الهدف (عربي):", type=["docx"], key="doc_ar")

    if doc_en and doc_ar:
        text_en, notes_en = extract_text_from_docx_bytes(doc_en.getvalue())
        text_ar, notes_ar = extract_text_from_docx_bytes(doc_ar.getvalue())

        if text_en.strip() and text_ar.strip():
            st.success("✅ تم استخراج النصوص من ملفات Word بنجاح وسرعة فائقة!")
            all_notes = [f"EN — {n}" for n in notes_en] + [f"AR — {n}" for n in notes_ar]
            if all_notes:
                with st.expander("ℹ️ ملاحظات استخراج مستندات Word"):
                    for n in all_notes:
                        st.write("•", n)

            segs_en = process_docx_text_to_segments(text_en)
            segs_ar = process_docx_text_to_segments(text_ar)

            aligned_initial = align_segments(segs_en, segs_ar)

            st.markdown("---")
            st.subheader("🧩 الخطوة 2: استوديو المحاذاة التفاعلي")
            
            initial_pairs = [{"en": e, "ar": a} for e, a in aligned_initial]
            token = hashlib.md5((text_en + "\x00" + text_ar).encode("utf-8")).hexdigest()

            studio_value = alignment_studio(pairs=initial_pairs, token=token, key="alignment_studio", default=None)
            current_pairs = studio_value if studio_value is not None else initial_pairs
            st.session_state["alignment_draft"] = current_pairs

            st.markdown("<br>", unsafe_allow_html=True)

            if st.button("🚀 اعتماد المحاذاة النهائية وتوليد TMX", use_container_width=True):
                current_pairs = st.session_state.get("alignment_draft", current_pairs)
                final_pairs = [
                    (p["en"].strip(), p["ar"].strip()) for p in current_pairs if p["en"].strip() and p["ar"].strip()
                ]
                if final_pairs:
                    st.session_state["aligned_pairs"] = final_pairs
                    st.session_state["tmx_bytes"] = generate_tmx_file(final_pairs)
                else:
                    st.session_state["tmx_bytes"] = None
                    st.error("⚠️ يرجى التأكد من وجود جمل صالحة للتقابل.")

            if st.session_state.get("tmx_bytes"):
                st.success(f"🎉 تم اعتماد **{len(st.session_state['aligned_pairs'])}** وحدة محاذاة بنجاح!")
                st.download_button(
                    "💾 تنزيل ملف الـ TMX الموحد",
                    data=st.session_state["tmx_bytes"],
                    file_name="Aligned_Legal_Memory.tmx",
                    mime="application/x-tmx+xml",
                )
        else:
            st.error("⚠️ تعذر استخراج النصوص من ملفات Word المرفوعة.")

# ---------------------------------------------------------
# Tab 2: الاستخراج الذكي وتوليد البطاقات
# ---------------------------------------------------------
with tab2:
    st.subheader("🔍 استخراج المصطلحات عبر الذكاء الاصطناعي")

    uploaded_tmx = st.file_uploader(
        "Upload your aligned .TMX memory file",
        type=["tmx"],
    )

    if uploaded_tmx is not None:
        aligned_units = parse_tmx_bytes(uploaded_tmx.read())
    else:
        aligned_units = st.session_state["aligned_pairs"]

    if aligned_units:
        if st.button("⚡ بدء الاستخراج الذكي والتصفية الصارمة"):
            if not api_key.strip():
                st.error("⚠️ يرجى إدخال مفتاح OpenRouter API Key أولاً.")
            else:
                client = OpenAI(
                    base_url="https://openrouter.ai/api/v1",
                    api_key=api_key.strip(),
                )

                system_prompt = """You are an expert legal terminologist specializing in general legal systems, international law, contracts, and legislations.

STRICT FILTERING & AUTOMATIC CONTEXT DETECTION RULES:

1. AUTOMATIC CONTEXT & SUBJECT FIELD DETECTION:
   - Carefully analyze the full text to deduce the highly specific "subject field" (e.g., International Public Law, Contractual Obligations, Commercial Law, Maritime Law).

2. POSITIVE SCOPE:
   - Extract ONLY substantive legal concepts, multiword legal collocations, obligations, rights, liabilities, statutory bodies, and judicial principles.

3. STRICT NEGATIVE DIRECTIVES:
   - DO NOT extract preamble, procedural, or introductory phrases (e.g., "Having reviewed", "Pursuant to", "In accordance with", "بعد الاطلاع على", "بناء على").
   - DO NOT extract dates, article numbers, or general administrative filler words.

4. OUTPUT FORMAT:
   - Output MUST be valid JSON as an object with key "terms", containing an array of objects EXACTLY matching this structure:
     {"ID": "", "en": "", "ar": "", "definition": "", "subject field": ""}
   - Ensure "definition" contains the Arabic legal explanation.
"""

                results = []
                progress_bar = st.progress(0)
                status_text = st.empty()

                def _s(item, key, default=""):
                    return str(item.get(key) or default).strip()

                MAX_CHARS = 8000 

                for i, (en_text, ar_text) in enumerate(aligned_units):
                    status_text.text(
                        f"جاري تحليل الوحدة ({i+1}/{len(aligned_units)})... عبر نموذج `{selected_model_id}`"
                    )

                    en_send = en_text[:MAX_CHARS // 2] if len(en_text) > MAX_CHARS // 2 else en_text
                    ar_send = ar_text[:MAX_CHARS // 2] if len(ar_text) > MAX_CHARS // 2 else ar_text
                    user_prompt = f"[EN]: {en_send}\n[AR]: {ar_send}"

                    try:
                        response = client.chat.completions.create(
                            model=selected_model_id,
                            messages=[
                                {"role": "system", "content": system_prompt},
                                {"role": "user", "content": user_prompt},
                            ],
                            temperature=0.0,
                            response_format={"type": "json_object"},
                        )

                        raw = json.loads(response.choices[0].message.content)
                        terms = raw.get("terms", []) if isinstance(raw, dict) else raw

                        if terms and isinstance(terms, list):
                            for item in terms:
                                if isinstance(item, dict) and item.get("en") and not _is_low_value_term(item.get("en")):
                                    results.append({
                                        "ID": _s(item, "ID", str(len(results) + 1)),
                                        "en": _s(item, "en"),
                                        "ar": _s(item, "ar"),
                                        "definition": _s(item, "definition"),
                                        "subject field": _s(item, "subject field", "قانون عام"),
                                        "Source EN": en_text,
                                        "Source AR": ar_text,
                                    })
                    except Exception as e:
                        st.warning(f"تنبيه في الوحدة {i+1}: {e}")

                    progress_bar.progress((i + 1) / len(aligned_units))

                status_text.text("✅ اكتمل الاستخراج والتصفية بنجاح!")

                if results:
                    df = pd.DataFrame(results)
                    df["_key"] = df["en"].map(_term_key) + "|" + df["ar"].map(_term_key)
                    df = df.groupby("_key", as_index=False, sort=False).agg({
                        "ID": "first", "en":"first", "ar":"first", "definition":"first", "subject field":"first",
                        "Source EN":lambda x: "\n---\n".join(dict.fromkeys(x)),
                        "Source AR":lambda x: "\n---\n".join(dict.fromkeys(x)),
                    })
                    
                    df.drop(columns=["_key"], inplace=True)
                    st.session_state["results_df"] = df

                    # التجهيز للإرسال إلى القاموس بالأعمدة الجديدة
                    terms_to_send = [
                        {
                            "ID": str(row.get("ID", "")).strip(),
                            "en": str(row.get("en", "")).strip(),
                            "ar": str(row.get("ar", "")).strip(),
                            "definition": str(row.get("definition", "")).strip(),
                            "subject field": str(row.get("subject field", "")).strip()
                        }
                        for _, row in df.iterrows() if row.get("en")
                    ]

                    if terms_to_send and send_to_dictionary(terms_to_send):
                        st.success("🚀 تم إرسال المصطلحات المستخرجة بنجاح إلى القاموس!")

                else:
                    st.error("لم يتم العثور على مصطلحات موضوعية بعد التصفية.")
                    
    # عرض النتائج والتصدير
    if st.session_state["results_df"] is not None:
        df_res = st.session_state["results_df"]

        st.subheader("📊 جدول المصطلحات المستخرجة:")
        edited_df = st.data_editor(
            df_res, use_container_width=True, hide_index=True, num_rows="dynamic",
            key="terms_editor",
        )
        st.session_state["results_df"] = edited_df
        df_res = edited_df

        # تجهيز التصدير ليقتصر على الأعمدة الخمسة فقط
        excel_buffer = io.BytesIO()
        excel_df = df_res[["ID", "en", "ar", "definition", "subject field"]]
        
        with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
            excel_df.to_excel(writer, index=False)
        excel_buffer.seek(0)

        try:
            pdf_buffer = generate_pdf_cards_bytes(df_res)
        except Exception as e:
            pdf_buffer = None
            st.warning(f"تعذر توليد بطاقات PDF: {e}")

        csv_data = excel_df.to_csv(index=False).encode("utf-8-sig")
        col_d1, col_d2, col_d3 = st.columns(3)
        with col_d1:
            st.download_button(
                label="📥 تنزيل مسرد المصطلحات (Excel)",
                data=excel_buffer,
                file_name="Extracted_Legal_Glossary.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        with col_d2:
            st.download_button("📄 تنزيل CSV (UTF-8)", data=csv_data, file_name="Extracted_Legal_Glossary.csv", mime="text/csv; charset=utf-8")
        with col_d3:
            if pdf_buffer is not None:
                st.download_button(
                    label="📑 تنزيل البطاقات الترجمية (PDF Cards)",
                    data=pdf_buffer,
                    file_name="Legal_Terminology_Cards.pdf",
                    mime="application/pdf",
                )

# =========================================================
# 10. التذييل بأسفل الصفحة
# =========================================================
st.markdown("<br><br>", unsafe_allow_html=True)
st.markdown(
    """
<div style="text-align: center; background-color: #131B2E; padding: 20px; border-radius: 12px; border: 1px solid #1E293B;">
    <h3 style="color: #EAB308; margin-bottom: 10px;">📫 تواصل معي (Contact & Support)</h3>
    <p style="font-size: 15px; color: #E5E7EB; margin: 5px;"><b>إعداد وتطوير:</b> يوسف نورالدين</p>
    <p style="font-size: 15px; color: #E5E7EB; margin: 5px;">📱 <b>الهاتف:</b> <a href="tel:0928756918" style="color: #EAB308; text-decoration: none;">0928756918</a> | ✉️ <b>البريد:</b> <a href="mailto:ywsfalmqmz@gmail.com" style="color: #EAB308; text-decoration: none;">ywsfalmqmz@gmail.com</a></p>
    <hr style="margin: 15px auto; width: 40%; border-top: 1px solid #1E293B;">
    <p style="font-size: 12px; color: #9CA3AF;">جميع الحقوق محفوظة © المنصة العامة للتحليل والاستخراج القانوني وتوليد البطاقات</p>
</div>
""",
    unsafe_allow_html=True,
)
